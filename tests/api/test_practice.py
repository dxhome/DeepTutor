"""Practice must preserve source evidence while imports/reviews are retry safe."""

import asyncio
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
import io
import json
import sqlite3
import time

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook
import pytest

from deeptutor.api.routers import practice, question_notebook
from deeptutor.core.assessment import ASSESSMENT_SOURCES
from deeptutor.services.practice.importing import normalize_question, preview
from deeptutor.services.practice.scheduler import DAY, day_bounds, schedule
from deeptutor.services.practice.storage import PracticeStore, ReviewConflict
from deeptutor.services.session.sqlite_store import SQLiteSessionStore


@pytest.fixture
def bank(tmp_path, monkeypatch):
    store = SQLiteSessionStore(tmp_path / "practice.db")
    for module in (practice, question_notebook):
        monkeypatch.setattr(module, "get_sqlite_session_store", lambda: store)
    app = FastAPI()
    app.include_router(question_notebook.router, prefix="/bank")
    app.include_router(practice.router, prefix="/practice")
    with TestClient(app) as client:
        yield store, PracticeStore(store.db_path), client


def seed(store, **changes):
    async def run():
        await store.ensure_notebook_session("lesson", "Lesson")
        item = {
            "question_id": "q1",
            "question": "2 + 2?",
            "question_type": "single_choice",
            "options": {"A": "3", "B": "4"},
            "correct_answer": "B",
            "user_answer": "A",
            "is_correct": False,
            **changes,
        }
        await store.upsert_notebook_entries("lesson", [item])
        return await store.find_notebook_entry("lesson", item["question_id"])

    return asyncio.run(run())


def test_document_import_mode_endpoint_reports_local_script(bank):
    _, _, client = bank
    response = client.get("/practice/import/document/mode")
    assert response.status_code == 200
    assert response.json() == {"mode": "local_script"}


@pytest.mark.parametrize("source", sorted(ASSESSMENT_SOURCES))
def test_every_producer_enrolls_mistakes_once(bank, source):
    store, repo, _ = bank
    entry = seed(store, source=source)
    first = repo.state(entry["id"])
    assert first["due_at"] == entry["created_at"] + DAY
    repo.review(entry["id"], "first_review", first["version"], "good", "B")
    reviewed = repo.state(entry["id"])
    seed(store, source=source)  # idempotent source re-save cannot reset the schedule
    assert repo.state(entry["id"]) == reviewed


def test_ungraded_imports_are_not_mistakes_and_review_preserves_original(bank):
    store, repo, client = bank
    entry = seed(store, result="ungraded")
    assert repo.overview("UTC")["mistakes"] == 0
    payload = {"request_id": "submission_00000001", "version": 0, "rating": "again", "answer": "A"}
    assert client.post(f"/practice/questions/{entry['id']}/review", json=payload).status_code == 200
    assert repo.overview("UTC")["mistakes"] == 1
    original = asyncio.run(store.get_notebook_entry(entry["id"]))
    for field in ("user_answer", "result", "is_correct", "correct_answer", "created_at"):
        assert original[field] == entry[field]
    assert client.get("/bank/entries?mistakes_only=true").json()["total"] == 1


def test_duplicate_review_and_two_tabs_cannot_advance_twice(bank):
    store, repo, _ = bank
    entry = seed(store)
    args = (entry["id"], "same_request", 0, "good", "B")
    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(lambda _: repo.review(*args), range(2)))
    assert results[0] == results[1]
    assert repo.state(entry["id"])["review_count"] == 1
    with pytest.raises(ReviewConflict):
        repo.review(entry["id"], "other_request", 0, "easy", "B")
    with pytest.raises(ReviewConflict):
        repo.review(entry["id"], "same_request", 0, "again", "A")


def test_server_rejects_incorrect_answer_with_success_rating(bank):
    store, repo, client = bank
    entry = seed(store)
    response = client.post(
        f"/practice/questions/{entry['id']}/review",
        json={"request_id": "submission_00000001", "version": 0, "rating": "good", "answer": "A"},
    )
    assert response.status_code == 422
    assert repo.state(entry["id"])["review_count"] == 0


def test_source_key_change_invalidates_open_question(bank):
    store, repo, _ = bank
    entry = seed(store)
    seed(store, correct_answer="A")
    with pytest.raises(ReviewConflict):
        repo.review(entry["id"], "stale", 0, "good", "A")


def test_timezone_due_order_resolved_and_empty_course_scope(bank):
    store, repo, client = bank
    old = seed(store)
    new = seed(store, question_id="q2")
    with repo.connect() as conn:
        conn.execute("UPDATE practice_review_state SET due_at=100 WHERE entry_id=?", (old["id"],))
        conn.execute("UPDATE practice_review_state SET due_at=200 WHERE entry_id=?", (new["id"],))
    assert repo.queue("Asia/Shanghai", now=DAY) == [old["id"], new["id"]]
    assert repo.queue("UTC", session_ids=[], now=DAY) == []
    asyncio.run(store.update_notebook_entry(old["id"], {"resolved": True}))
    assert repo.queue("UTC", now=DAY) == [new["id"]]
    assert repo.overview("UTC", now=DAY)["mistakes"] == 2
    with repo.connect() as conn:
        conn.execute("UPDATE sessions SET deleted_at=1 WHERE id='lesson'")
    assert repo.overview("UTC", now=DAY)["total"] == 0
    assert repo.queue("UTC", now=DAY) == []
    assert client.get(f"/practice/questions/{old['id']}").status_code == 404
    assert (
        client.post(f"/practice/questions/{old['id']}/check", json={"answer": "B"}).status_code
        == 404
    )


def test_first_mistake_timestamp_follows_first_failed_practice(bank):
    store, repo, _ = bank
    entry = seed(store, result="ungraded")
    repo.review(entry["id"], "correct_first", 0, "good", "B", now=100)
    repo.review(entry["id"], "wrong_later", 1, "again", "A", now=200)
    assert repo.state(entry["id"])["first_wrong_at"] == 200


def test_source_mistake_after_success_uses_the_actual_wrong_date(bank):
    store, repo, _ = bank
    entry = seed(store, result="ungraded")
    repo.review(entry["id"], "initial-success", 0, "good", "B", now=100)
    with repo.connect() as conn:
        conn.execute(
            "UPDATE notebook_entries SET result='incorrect', is_correct=0, updated_at=200 WHERE id=?",
            (entry["id"],),
        )
    assert repo.state(entry["id"])["first_wrong_at"] == 200


def test_oversized_spreadsheet_dimensions_and_csv_cells_are_rejected(bank):
    _, _, client = bank
    book = Workbook()
    book.active.cell(row=1, column=1000, value="question")
    output = io.BytesIO()
    book.save(output)
    book.close()
    for name, data in [
        ("wide.xlsx", output.getvalue()),
        ("large.csv", b"question,answer\n" + b"a" * 200000 + b",yes\n"),
    ]:
        response = client.post("/practice/import/preview", files={"file": (name, data)})
        assert response.status_code == 422


def test_import_preview_commit_retries_dedup_tags_and_no_source_overwrite(bank):
    store, repo, client = bank
    data = "题目,题型,A,B,答案,标签\n2+2=?,单选题,3,4,B,数学\n".encode()
    result = client.post(
        "/practice/import/preview", files={"file": ("questions.csv", data)}, data={"target": "bank"}
    ).json()
    assert result["valid"] == 1 and result["token"]
    assert repo.overview("UTC")["total"] == 0  # preview does not alter the bank
    commit = client.post("/practice/import/commit", json={"token": result["token"]})
    assert commit.json() == {"created": 1, "duplicates": 0, "total": 1}
    assert (
        client.post("/practice/import/commit", json={"token": result["token"]}).json()
        == commit.json()
    )
    questions = [
        normalize_question({"question": "2+2=?", "A": "3", "B": "4", "answer": "B", "tags": "复习"})
    ]
    again = repo.commit_import(repo.stage_import("again.json", "mistakes", questions))
    assert again["duplicates"] == 1
    listing = asyncio.run(store.list_notebook_entries())
    assert listing["total"] == 1
    entry = listing["items"][0]
    assert entry["result"] == "ungraded"
    assert {category["name"] for category in entry["categories"]} == {"数学", "复习"}
    assert repo.overview("UTC")["mistakes"] == 1


def test_invalid_row_blocks_entire_import_and_reports_line(bank):
    _, repo, client = bank
    result = client.post(
        "/practice/import/preview",
        files={"file": ("file.csv", b"question,answer\ngood,yes\nbad,\n")},
    ).json()
    assert result["token"] is None
    assert result["errors"][0]["row"] == 3
    assert result["valid"] == 1
    assert repo.overview("UTC")["total"] == 0


def test_template_round_trips_and_formulas_are_rejected(bank):
    _, _, client = bank
    for format in ("csv", "xlsx"):
        data = client.get(f"/practice/import/template?format={format}").content
        assert len(preview(data, f"template.{format}")["questions"]) == 4
    book = Workbook()
    book.active.append(["question", "answer"])
    book.active.append(["=1+1", "2"])
    file = io.BytesIO()
    book.save(file)
    with pytest.raises(ValueError, match="formulas"):
        preview(file.getvalue(), "formulas.xlsx")


def test_multi_select_text_and_json_validation():
    row = normalize_question(
        {"题干": "Primes?", "题型": "多选题", "options": ["2", "3", "4"], "答案": "AB"}
    )
    assert row["correct_answer"] == "A,B"
    assert preview(json.dumps([row]).encode(), "q.json")["questions"][0]["question"] == "Primes?"
    with pytest.raises(ValueError, match="500"):
        preview(json.dumps([row] * 501).encode(), "many.json")


def test_additive_migration_is_idempotent(bank):
    store, repo, _ = bank
    entry = seed(store)
    with repo.connect() as conn:
        for trigger in (
            "practice_capture_insert",
            "practice_capture_update",
            "practice_invalidate_question",
        ):
            conn.execute(f"DROP TRIGGER {trigger}")
        conn.execute("DROP TABLE practice_review_state")
    upgraded = SQLiteSessionStore(store.db_path)
    state = repo.state(entry["id"])
    assert state["first_wrong_at"] == entry["created_at"]
    restored = asyncio.run(upgraded.get_notebook_entry(entry["id"]))
    assert {k: v for k, v in restored.items() if k not in {"practice", "categories"}} == entry
    SQLiteSessionStore(store.db_path)
    assert repo.state(entry["id"]) == state
    asyncio.run(upgraded.delete_notebook_entry(entry["id"]))
    with repo.connect() as conn:
        assert conn.execute("SELECT COUNT(*) FROM practice_review_state").fetchone()[0] == 0


def test_import_transaction_rolls_back(bank):
    _, repo, _ = bank
    token = repo.stage_import(
        "q.json", "bank", [normalize_question({"question": "Q", "answer": "A"})]
    )
    with repo.connect() as conn:
        conn.execute(
            "CREATE TRIGGER reject_import BEFORE INSERT ON notebook_entries BEGIN SELECT RAISE(ABORT, 'unavailable'); END"
        )
    with pytest.raises(sqlite3.IntegrityError):
        repo.commit_import(token)
    assert repo.overview("UTC")["total"] == 0
    with repo.connect() as conn:
        conn.execute("DROP TRIGGER reject_import")
    assert repo.commit_import(token)["created"] == 1


def test_daily_scheduler_and_dst():
    now = datetime(2026, 3, 8, 12, tzinfo=timezone.utc).timestamp()
    start, end = day_bounds("America/New_York", now)
    assert end - start == 23 * 3600
    with pytest.raises(ValueError):
        day_bounds("not/a-zone", now)
    good = schedule({}, "good", now)
    better = schedule(good, "good", good["due_at"])
    again = schedule(better, "again", better["due_at"])
    assert good["interval_days"] < better["interval_days"]
    assert 0 < again["interval_days"] < 1
    assert again["lapses"] == 1 and again["review_count"] == 3
    assert schedule({}, "easy", now)["due_at"] > good["due_at"]
    hard = schedule(better, "hard", now)
    assert (
        schedule(hard, "good", now)["interval_days"] >= schedule(hard, "hard", now)["interval_days"]
    )


def test_malformed_files_are_validation_errors(bank):
    _, _, client = bank
    for filename, data in [
        ("broken.xlsx", b"not a workbook"),
        ("broken.json", b"{"),
        ("old.xls", b"old format"),
    ]:
        response = client.post("/practice/import/preview", files={"file": (filename, data)})
        assert response.status_code == 422


def test_document_import_requires_review_before_committing(bank, monkeypatch):
    _, repo, client = bank
    from deeptutor.services.practice.document_import import validate_candidate

    candidate = {
        "id": "1", "number": "1", "page": 1, "source_excerpt": "1. Compute 2+2",
        "question": "Compute 2+2", "question_type": "fill_blank", "options": {},
        "correct_answer": "", "explanation": "", "tags": [],
        "warnings": ["Reference answer missing"], "confirmed": False, "selected": True,
    }
    candidate["errors"] = validate_candidate(candidate)

    async def fake_parse(_data, _filename, **_kwargs):
        return [candidate]

    monkeypatch.setattr(practice, "parse_document", fake_parse)
    response = client.post(
        "/practice/import/document/preview",
        files={"file": ("exam.md", b"1. Compute 2+2")},
    )
    assert response.status_code == 200
    draft = response.json()
    assert draft["items"][0]["errors"]
    assert repo.overview("UTC")["total"] == 0
    assert client.post("/practice/import/commit", json={"token": draft["token"]}).status_code == 422

    item = {**draft["items"][0], "correct_answer": "4", "confirmed": True}
    saved = client.patch(
        f"/practice/import/document/{draft['token']}",
        json={"revision": draft["revision"], "items": [item], "target": "mistakes"},
    ).json()
    assert saved["revision"] == 1 and saved["target"] == "mistakes"
    assert saved["status"] == "ready" and not saved["items"][0]["errors"]
    assert client.patch(
        f"/practice/import/document/{draft['token']}",
        json={"revision": draft["revision"], "items": [item], "target": "bank"},
    ).status_code == 409
    committed = client.post("/practice/import/commit", json={"token": draft["token"]})
    assert committed.json() == {"created": 1, "duplicates": 0, "total": 1}
    assert repo.overview("UTC")["mistakes"] == 1
    assert client.post("/practice/import/commit", json={"token": draft["token"]}).json() == committed.json()
    assert client.get(f"/practice/import/document/{draft['token']}/source").content == b"1. Compute 2+2"


def test_markdown_document_uses_numbering_and_matching_answers_without_models():
    from deeptutor.services.practice.document_import import parse_document

    text = """一、填空题（共2题）
1. 2 + 2 = ____
2. 3 + 3 = ____
参考答案
1.【答案】4
2.【答案】6
"""
    result = asyncio.run(parse_document(text.encode(), "paper.md"))
    assert [(item["number"], item["correct_answer"], item["question_type"]) for item in result] == [
        ("1", "4", "fill_blank"), ("2", "6", "fill_blank"),
    ]
    assert all(not item["errors"] for item in result)
    assert "题目原文（第 1 页）" in result[0]["source_excerpt"]
    assert "答案提取" in result[0]["source_excerpt"]
    assert "2 + 2 = ____" in result[0]["source_excerpt"]
    assert "3 + 3 = ____" not in result[0]["source_excerpt"]


def test_answer_number_gaps_do_not_attach_later_answers_to_previous_question():
    from deeptutor.services.practice.document_import import parse_document

    text = """一、填空题（共3题）
1. First ____
2. Second ____
3. Third ____
参考答案
1.【答案】A
3.【答案】C
"""
    result = asyncio.run(parse_document(text.encode(), "paper.md"))
    assert [item["correct_answer"] for item in result] == ["A", "", "C"]
    assert any("未找到" in warning for warning in result[1]["warnings"])


def test_section_heading_closes_previous_question_and_is_not_added_to_its_stem():
    from deeptutor.services.practice.document_import import parse_document

    text = """一、填空题（共1题）
1. Fill this blank ____
二、单选题（共1题）每题只有一个正确选项。
2. Pick one.
A. First B. Second
参考答案
1.【答案】filled
2.【答案】A
"""
    result = asyncio.run(parse_document(text.encode(), "paper.md"))
    assert len(result) == 2
    assert "正确选项" not in result[0]["question"]
    assert result[1]["question_type"] == "single_choice"


def test_spaced_answer_labels_and_answer_section_headings_do_not_leak_between_answers():
    from deeptutor.services.practice.document_import import parse_document

    text = """一、填空题（共2题）
1. First ____
2. Second ____
参考答案
1.【 答 案 】A
【 解 析 】Only explanation for the first item.
二、单选题
2.【 答 案 】B
"""
    result = asyncio.run(parse_document(text.encode(), "paper.md"))
    assert [item["correct_answer"] for item in result] == ["A", "B"]
    assert "Only explanation for the first item." in result[0]["explanation"]
    assert "答案提取（答案区第 1 页）:\nA" in result[0]["source_excerpt"]
    assert "答案提取（答案区第 1 页）:\nB" not in result[0]["source_excerpt"]


def test_script_parser_flags_unmatched_answers_instead_of_guessing():
    from deeptutor.services.practice.document_import import parse_document

    result = asyncio.run(parse_document(
        "1. Find x\n参考答案\n".encode(), "paper.md",
    ))
    assert result[0]["correct_answer"] == ""
    assert result[0]["errors"]
    assert any("未找到" in warning for warning in result[0]["warnings"])


def test_script_parser_classifies_choice_options_and_answer():
    from deeptutor.services.practice.document_import import parse_document

    text = """一、单选题（共1题）
1. 2 + 2 = ?
A. 3 B. 4 C. 5 D. 6
参考答案
1.【答案】B
"""
    item = asyncio.run(parse_document(text.encode(), "paper.md"))[0]
    assert item["question_type"] == "single_choice"
    assert item["options"] == {"A": "3", "B": "4", "C": "5", "D": "6"}
    assert item["correct_answer"] == "B"
    assert item["errors"] == []


def test_word_document_is_parsed_locally():
    from docx import Document
    from deeptutor.services.practice.document_import import parse_document

    document = Document()
    document.add_paragraph("一、填空题（共1题）")
    document.add_paragraph("1. What is 2+2? ____")
    document.add_paragraph("参考答案")
    document.add_paragraph("1.【答案】4")
    buffer = io.BytesIO()
    document.save(buffer)
    result = asyncio.run(parse_document(buffer.getvalue(), "paper.docx"))
    assert result[0]["question"] == "What is 2+2? ____"
    assert result[0]["correct_answer"] == "4"
    assert result[0]["errors"] == []


def test_scanned_pdf_fails_with_clear_no_ocr_message():
    import pymupdf
    from deeptutor.services.practice.document_import import parse_document

    document = pymupdf.open()
    document.new_page()
    data = document.tobytes()
    document.close()
    with pytest.raises(ValueError, match="scanned PDFs require OCR"):
        asyncio.run(parse_document(data, "scan.pdf"))


def test_corrupt_word_document_reports_a_repairable_error():
    from deeptutor.services.practice.document_import_rules import _read_document_pages

    with pytest.raises(ValueError, match="docx|Word|archive|valid|read"):
        _read_document_pages(b"not a Word archive", ".docx")


def test_document_preview_marks_duplicates_within_the_same_draft(bank):
    _, repo, _ = bank
    first = {
        "id": "1", "number": "1", "page": 1, "source_excerpt": "Q",
        "question": "Q", "question_type": "fill_blank", "options": {},
        "correct_answer": "A", "explanation": "", "tags": [],
        "warnings": [], "errors": [], "confirmed": False, "selected": True,
    }
    token = repo.stage_document(
        "paper.md", "bank", [first, {**first, "id": "2", "number": "2"}],
        b"Q", "text/markdown",
    )
    assert [item["duplicate"] for item in repo.get_document(token)["items"]] == [False, True]
    assert repo.commit_import(token) == {"created": 1, "duplicates": 1, "total": 2}


def test_document_job_can_be_resumed_after_upload(bank, monkeypatch):
    _, _, client = bank

    async def fake_parse(_data, _filename, **kwargs):
        await kwargs["progress"]({"stage": "extracting_questions", "percent": 70,
                                   "message": "Extracted one question"})
        return [{
            "id": "1", "number": "1", "page": 1, "source_excerpt": "Q",
            "question": "Q", "question_type": "fill_blank", "options": {},
            "correct_answer": "A", "explanation": "", "tags": [], "warnings": [],
            "errors": [], "confirmed": False, "selected": True,
        }]

    monkeypatch.setattr(practice, "parse_document", fake_parse)
    started = client.post(
        "/practice/import/document/start",
        files={"file": ("exam.md", b"Q")},
    )
    assert started.status_code == 200
    token = started.json()["token"]
    for _ in range(30):
        resumed = client.get(f"/practice/import/document/{token}").json()
        if resumed["status"] == "ready":
            break
        time.sleep(0.02)
    assert resumed["status"] == "ready"
    assert resumed["percent"] == 100
    assert "Extracted one question" in resumed["logs"]
    assert resumed["items"][0]["question"] == "Q"
    listed = client.get("/practice/import/document/jobs").json()
    assert listed[0]["filename"] == "exam.md" and listed[0]["percent"] == 100


def test_document_job_timeout_fails_and_can_be_retried(bank, monkeypatch):
    _, _, client = bank

    async def slow_parse(_data, _filename, **_kwargs):
        await asyncio.sleep(1)
        return []

    monkeypatch.setattr(practice, "parse_document", slow_parse)
    monkeypatch.setattr(practice, "DOCUMENT_JOB_TIMEOUT", 0.02)
    started = client.post("/practice/import/document/start", files={"file": ("slow.md", b"Q")}).json()
    token = started["token"]
    for _ in range(40):
        draft = client.get(f"/practice/import/document/{token}").json()
        if draft["status"] == "failed":
            break
        time.sleep(0.02)
    assert draft["status"] == "failed"
    assert "timed out" in draft["error"]
    assert "Failed:" in draft["logs"][-1]


def test_queued_document_can_be_cancelled_without_starting(bank):
    _, repo, _ = bank
    token = repo.stage_document_job("queued.md", "bank", b"Q", "text/markdown")
    assert repo.get_document(token)["status"] == "queued"
    repo.cancel_document_job(token)
    assert repo.start_document_job(token) is False
    repo.finish_document_job(token, [])
    draft = repo.get_document(token)
    assert draft["status"] == "cancelled"
    assert draft["logs"][-1] == "Parsing cancelled"


@pytest.mark.parametrize("status", ["queued", "processing", "failed", "ready"])
def test_uncommitted_document_jobs_can_be_deleted(bank, status):
    _, repo, client = bank
    token = repo.stage_document_job(f"{status}.md", "bank", b"Q", "text/markdown")
    if status in {"processing", "failed", "ready"}:
        repo.start_document_job(token)
    if status == "failed":
        repo.finish_document_job(token, [], "parse failed")
    elif status == "ready":
        repo.finish_document_job(token, [])
    response = client.delete(f"/practice/import/document/{token}")
    assert response.status_code == 200
    assert response.json() == {"status": "deleted"}
    assert client.get(f"/practice/import/document/{token}").status_code == 404


def test_stalled_document_job_reports_failure_after_restart(bank):
    _, repo, _ = bank
    token = repo.stage_document_job("stalled.md", "bank", b"Q", "text/markdown")
    with repo.connect() as conn:
        conn.execute("UPDATE practice_import_jobs SET updated_at=? WHERE token=?",
                     (time.time() - 1900, token))
    draft = repo.get_document(token)
    assert draft["status"] == "failed"
    assert "retry" in draft["error"]


def test_failed_document_job_can_retry_without_reupload(bank, monkeypatch):
    _, _, client = bank
    calls = 0

    async def sometimes_parse(_data, _filename, **_kwargs):
        nonlocal calls
        calls += 1
        if calls == 1:
            raise RuntimeError("temporary model outage")
        return [{
            "id": "1", "number": "1", "page": 1, "source_excerpt": "Q",
            "question": "Q", "question_type": "fill_blank", "options": {},
            "correct_answer": "A", "explanation": "", "tags": [], "warnings": [],
            "errors": [], "confirmed": False, "selected": True,
        }]

    monkeypatch.setattr(practice, "parse_document", sometimes_parse)
    started = client.post(
        "/practice/import/document/start", files={"file": ("exam.md", b"Q")}
    ).json()
    token = started["token"]
    for _ in range(30):
        status = client.get(f"/practice/import/document/{token}").json()["status"]
        if status == "failed":
            break
        time.sleep(0.02)
    assert status == "failed"
    assert client.post(f"/practice/import/document/{token}/retry").status_code == 200
    for _ in range(30):
        status = client.get(f"/practice/import/document/{token}").json()["status"]
        if status == "ready":
            break
        time.sleep(0.02)
    assert status == "ready" and calls == 2


def test_import_receipts_and_questions_are_isolated_between_stores(bank, tmp_path):
    _, repo, _ = bank
    token = repo.stage_import(
        "q.json", "bank", [normalize_question({"question": "Q", "answer": "A"})]
    )
    other = SQLiteSessionStore(tmp_path / "other-user.db")
    other_repo = PracticeStore(other.db_path)
    with pytest.raises(LookupError):
        other_repo.commit_import(token)
    repo.commit_import(token)
    assert other_repo.overview("UTC")["total"] == 0
    with pytest.raises(LookupError):
        other_repo.review(1, "other-user-attempt", 0, "again", "")


def test_course_imports_keep_their_course_and_file_provenance(bank, monkeypatch):
    from types import SimpleNamespace

    monkeypatch.setattr(
        "deeptutor.services.courses.get_course_service", lambda: SimpleNamespace(get=lambda _: {})
    )
    store, repo, client = bank
    question = normalize_question({"question": "Q", "answer": "A"})
    repo.commit_import(repo.stage_import("lesson.csv", "mistakes", [question], "course-a"))
    listing = asyncio.run(store.list_notebook_entries())
    assert listing["total"] == 1
    item = listing["items"][0]
    assert item["session_id"] == ""
    assert item["origin_type"] == "external_import"
    assert item["origin_ref"] == "practice-import"
    assert item["material_id"].startswith("import:")
    assert repo.overview("UTC")["mistakes"] == 1
    # Import provenance must not manufacture a conversation or leak into a
    # conversation-derived course scope.
    assert asyncio.run(store.list_sessions()) == []
    assert client.get("/bank/entries?course_id=course-a").json()["total"] == 0
    assert client.get("/practice/summary?course_id=course-a").json()["mistakes"] == 0
    assert client.get("/bank/entries?course_id=course-b").json()["total"] == 0


def test_analytics_separates_creation_first_mistake_and_repeat_reviews(bank):
    store, repo, client = bank
    now = datetime(2026, 9, 19, 12, tzinfo=timezone.utc).timestamp()
    entry = seed(store, source="mastery_path", result="ungraded")
    with repo.connect() as conn:
        conn.execute(
            "UPDATE notebook_entries SET created_at=? WHERE id=?", (now - 3 * DAY, entry["id"])
        )
    repo.review(entry["id"], "first-wrong", 0, "again", "A", now=now - DAY)
    repo.review(entry["id"], "second-wrong", 1, "again", "A", now=now)
    repo.review(entry["id"], "second-wrong", 1, "again", "A", now=now)
    report = repo.analytics("Asia/Shanghai", 7, now=now)
    assert report["totals"] == {"questions": 1, "mistakes": 1, "reviews": 2}
    assert len(report["daily"]) == 7
    assert report["daily"][-4] == {
        "date": "2026-09-16",
        "questions": 1,
        "mistakes": 0,
        "reviews": 0,
    }
    assert report["daily"][-2]["mistakes"] == 1
    assert report["daily"][-1]["mistakes"] == 0
    assert report["sources"] == [{"source": "mastery_path", **report["totals"]}]
    asyncio.run(store.update_notebook_entry(entry["id"], {"resolved": True}))
    assert repo.analytics("Asia/Shanghai", 7, now=now)["totals"] == report["totals"]
    assert repo.analytics("Asia/Shanghai", 7, session_ids=[], now=now)["totals"]["questions"] == 0
    assert client.get("/practice/analytics?days=8").status_code == 422
    assert client.get("/practice/analytics?timezone=not-a-zone").status_code == 422


def test_analytics_local_midnight_dst_zero_days_and_deleted_sources(bank):
    store, repo, _ = bank
    now = datetime(2026, 3, 10, 12, tzinfo=timezone.utc).timestamp()
    # These straddle local midnight during the daylight-saving transition.
    entries = [seed(store, question_id=f"dst-{i}", source="book") for i in range(2)]
    timestamps = [
        datetime(2026, 3, 8, 4, 59, tzinfo=timezone.utc).timestamp(),
        datetime(2026, 3, 8, 5, 1, tzinfo=timezone.utc).timestamp(),
    ]
    with repo.connect() as conn:
        for entry, ts in zip(entries, timestamps, strict=True):
            conn.execute("UPDATE notebook_entries SET created_at=? WHERE id=?", (ts, entry["id"]))
            conn.execute(
                "UPDATE practice_review_state SET first_wrong_at=? WHERE entry_id=?",
                (ts, entry["id"]),
            )
    data = repo.analytics("America/New_York", 7, now=now)
    by_day = {row["date"]: row for row in data["daily"]}
    assert by_day["2026-03-07"]["questions"] == by_day["2026-03-08"]["questions"] == 1
    assert by_day["2026-03-09"]["questions"] == 0
    with repo.connect() as conn:
        conn.execute("UPDATE sessions SET deleted_at=? WHERE id='lesson'", (now,))
    assert repo.analytics("America/New_York", 7, now=now)["sources"] == []
